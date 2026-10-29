import os
import uuid
from io import BytesIO
from pathlib import Path
from fastapi import APIRouter, Depends, HTTPException, Query, Response, UploadFile, File, Form
from sqlalchemy.orm import Session
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.utils.lkh_status import get_status_pengisian

from app.db.database import get_db
from app.models.rkh import RKH
from app.models.lkh import LKH
from app.models.user import User
from app.models.customer import Customer
from app.models.project_prospek import ProjectProspek
from app.models.project_status_history import ProjectStatusHistory
from app.models.project_status import ProjectStatus
from app.schemas.lkh import LKHCreate, LKHUpdate, LKHResponse,LKHExportSchema
from app.core.dependencies import get_current_user, require_roles
from app.utils.project_log import catat_project_log
from app.core.roles import LKH_VIEW_ROLES
from datetime import date, datetime, timezone
from zoneinfo import ZoneInfo

router = APIRouter(prefix="/lkh", tags=["LKH"])

WIB = ZoneInfo("Asia/Jakarta")

STATUS_LKH_LABELS = {
    "NOT_YET": "Open",
    "DONE": "Done",
}

STATUS_LKH_STYLES = {
    "NOT_YET": ("FFF2CC", "7F6000"),
    "DONE": ("E2EFDA", "375623"),
}


def to_wib(value):
    if value is None:
        return None
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.astimezone(WIB).replace(tzinfo=None)


def build_lkh_workbook(rows, filter_text, exported_by, exported_at):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Data LKH"

    headers = [
        "Tanggal Kunjungan",
        "Sales",
        "Customer",
        "Project",
        "Kegiatan",
        "Nama PIC",
        "Nomor PIC",
        "Hasil Kunjungan",
        "Tindak Lanjut",
        "Tanggal Follow Up",
        "Status LKH",
        "Latitude",
        "Longitude",
        "Dibuat Pada",
        "Diperbarui Pada",
    ]

    worksheet.append(headers)

    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78",
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    for item in rows:
        lkh = item["lkh"]
        status = lkh.status_lkh

        worksheet.append([
            lkh.tanggal_kunjungan,
            item["sales"],
            item["customer"],
            item["project"],
            item["kegiatan"],
            lkh.nama_pic,
            str(lkh.nomor_pic) if lkh.nomor_pic is not None else None,
            lkh.hasil_kunjungan,
            lkh.tindak_lanjut,
            lkh.tanggal_follow_up,
            STATUS_LKH_LABELS.get(status, status),
            float(lkh.latitude) if lkh.latitude is not None else None,
            float(lkh.longitude) if lkh.longitude is not None else None,
            to_wib(lkh.created_at),
            to_wib(lkh.updated_at),
        ])

    for row in worksheet.iter_rows(min_row=2):
        row[0].number_format = "dd-mm-yyyy"
        row[6].number_format = "@"
        row[9].number_format = "dd-mm-yyyy"
        row[11].number_format = "0.000000"
        row[12].number_format = "0.000000"
        row[13].number_format = "dd-mm-yyyy hh:mm:ss"
        row[14].number_format = "dd-mm-yyyy hh:mm:ss"

    for column_cells in worksheet.columns:
        max_length = max(
            len(str(cell.value or ""))
            for cell in column_cells
        )
        column_letter = get_column_letter(
            column_cells[0].column
        )
        worksheet.column_dimensions[
            column_letter
        ].width = min(max_length + 3, 40)

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    return workbook


def get_lkh_query(db: Session, current_user: User):
    query = db.query(LKH)

    if current_user.role == "SALES":
        query = query.filter(LKH.sales_id == current_user.id)
    elif current_user.role not in LKH_VIEW_ROLES:
        raise HTTPException(
            status_code=403,
            detail="Kamu tidak memiliki akses ke laporan kunjungan",
        )
    return query


@router.get("/", response_model=list[LKHResponse])
def get_lkh(
    rkh_id: int | None = None,
    sales_id: int | None = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    query = get_lkh_query(db, current_user)
    if rkh_id is not None:
        query = query.filter(LKH.rkh_id == rkh_id)
    if sales_id is not None:
        if current_user.role == "SALES" and sales_id != current_user.id:
            raise HTTPException(
                status_code=403,
                detail="Kamu tidak memiliki akses ke laporan sales lain",
            )
        query = query.filter(LKH.sales_id == sales_id)

    return query.order_by(LKH.id.desc()).all()




@router.get("/export/excel")
def export_lkh_excel(
    sales_id: list[int] | None = Query(None),
    customer_id: list[int] | None = Query(None),
    tanggal_mulai: date | None = None,
    tanggal_selesai: date | None = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    query = get_lkh_query(db, current_user)

    if (
        tanggal_mulai is not None
        and tanggal_selesai is not None
        and tanggal_mulai > tanggal_selesai
    ):
        raise HTTPException(
            status_code=400,
            detail="Tanggal mulai tidak boleh melebihi tanggal selesai",
        )

    if tanggal_mulai is not None:
        query = query.filter(LKH.tanggal_kunjungan >= tanggal_mulai)

    if tanggal_selesai is not None:
        query = query.filter(LKH.tanggal_kunjungan <= tanggal_selesai)

    sales_ids = sorted(set(sales_id or []))
    if sales_ids:
        query = query.filter(LKH.sales_id.in_(sales_ids))

    customer_ids = sorted(set(customer_id or []))
    if customer_ids:
        query = query.join(RKH, RKH.id == LKH.rkh_id).filter(
            RKH.customer_id.in_(customer_ids)
        )

    data_lkh = query.order_by(LKH.tanggal_kunjungan.desc(), LKH.id.desc()).all()

    rkh_map = {}
    customer_map = {}
    project_map = {}
    user_map = {}

    if data_lkh:
        rkh_map = {
            item.id: item
            for item in db.query(RKH).filter(
                RKH.id.in_({lkh.rkh_id for lkh in data_lkh})
            ).all()
        }
        customer_ids_in_data = {rkh.customer_id for rkh in rkh_map.values() if rkh.customer_id}
        project_ids = {rkh.project_id for rkh in rkh_map.values() if rkh.project_id}
        sales_ids_in_data = {lkh.sales_id for lkh in data_lkh if lkh.sales_id}

        if customer_ids_in_data:
            customer_map = {
                item.id: item.nama_instansi
                for item in db.query(Customer).filter(Customer.id.in_(customer_ids_in_data)).all()
            }
        if project_ids:
            project_map = {
                item.id: item.nama_project
                for item in db.query(ProjectProspek).filter(ProjectProspek.id.in_(project_ids)).all()
            }
        if sales_ids_in_data:
            user_map = {
                item.id: item.nama
                for item in db.query(User).filter(User.id.in_(sales_ids_in_data)).all()
            }

    rows = []
    for lkh in data_lkh:
        rkh = rkh_map.get(lkh.rkh_id)
        rows.append(
            {
                "lkh": lkh,
                "sales": user_map.get(lkh.sales_id),
                "customer": customer_map.get(rkh.customer_id) if rkh else None,
                "project": project_map.get(rkh.project_id) if rkh else None,
                "kegiatan": rkh.judul_kegiatan if rkh else None,
            }
        )

    if tanggal_mulai or tanggal_selesai:
        mulai_text = tanggal_mulai.strftime("%d-%m-%Y") if tanggal_mulai else "..."
        selesai_text = tanggal_selesai.strftime("%d-%m-%Y") if tanggal_selesai else "..."
        periode_text = f"{mulai_text} s/d {selesai_text}"
    else:
        periode_text = "Semua"

    sales_text = "Semua"
    if sales_ids:
        names = [
            name
            for (name,) in db.query(User.nama)
            .filter(User.id.in_(sales_ids))
            .order_by(User.nama.asc())
            .all()
        ]
        sales_text = ", ".join(names) or "Semua"

    customer_text = "Semua"
    if customer_ids:
        names = [
            name
            for (name,) in db.query(Customer.nama_instansi)
            .filter(Customer.id.in_(customer_ids))
            .order_by(Customer.nama_instansi.asc())
            .all()
        ]
        customer_text = ", ".join(names) or "Semua"

    filter_text = (
        f"Sales: {sales_text} | "
        f"Customer: {customer_text} | "
        f"Periode kunjungan: {periode_text}"
    )

    now = datetime.now(WIB)
    workbook = build_lkh_workbook(rows, filter_text, current_user.nama, now)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    filename = f"Data_LKH_{now.strftime('%d-%m-%Y_%H%M')}.xlsx"

    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"'
        },
    )


@router.get("/{lkh_id}", response_model=LKHResponse)
def get_lkh_detail(
    lkh_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    laporan = get_lkh_query(db, current_user).filter(
        LKH.id == lkh_id
    ).first()
    if not laporan:
        raise HTTPException(
            404,
            "Laporan tidak ditemukan atau kamu tidak memiliki akses",
        )

    return laporan


@router.post("/", response_model=LKHResponse, status_code=201)
def create_lkh(
    rkh_id: int = Form(...),
    tanggal_kunjungan: date = Form(...),
    hasil_kunjungan: str = Form(...),
    nama_pic: str = Form(...),
    nomor_pic: str = Form(...),
    tindak_lanjut: str | None = Form(None),
    tanggal_follow_up: date | None = Form(None),
    latitude: float | None = Form(None),
    longitude: float | None = Form(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(
        require_roles("SUPERADMIN", "SPV", "SALES")
    ),
):
    rkh = db.query(RKH).filter(RKH.id == rkh_id).first()

    if not rkh:
        raise HTTPException(404, "RKH tidak ditemukan")

    if current_user.role == "SALES" and rkh.sales_id != current_user.id:
            raise HTTPException(403, "RKH ini bukan tanggung jawabmu")

        # Cek status keterlambatan pengisian LKH
    status_pengisian = get_status_pengisian(rkh)

    if (
        status_pengisian == "LOCKED_LATE"
        and current_user.role != "SPV"
    ):
        raise HTTPException(
            status_code=403,
            detail={
                "status": "LOCKED_LATE",
                "message": (
                    "LKH sudah terkunci. "
                    "Hanya SPV yang dapat mengisi LKH "
                    "lebih dari satu hari setelah tanggal rencana."
                ),
            },
        )

    if existing:
        raise HTTPException(409, "RKH ini sudah memiliki laporan kunjungan")

    try:
        laporan = LKH(
            rkh_id=rkh.id,
            sales_id=rkh.sales_id,
            tanggal_kunjungan=tanggal_kunjungan,
            hasil_kunjungan=hasil_kunjungan,
            nama_pic=nama_pic,
            nomor_pic=nomor_pic,
            tindak_lanjut=tindak_lanjut,
            tanggal_follow_up=tanggal_follow_up,
            latitude=latitude,
            longitude=longitude,
            status_pengisian=status_pengisian,
        )

        db.add(laporan)
        db.flush()
        project = (
            db.query(ProjectProspek)
            .filter(ProjectProspek.id == rkh.project_id)
            .first()
        )
        if project:
            catat_project_log(
                db=db,
                project_id=project.id,
                user_id=current_user.id,
                aktivitas="Laporan kunjungan dibuat",
                deskripsi=(
                    f"Laporan kunjungan untuk RKH tanggal "
                    f"{rkh.tanggal_rencana} telah dibuat."
                ),
            )
        rkh.status = "CLOSED"
        db.commit()
        db.refresh(laporan)

        return laporan

    except Exception:
        db.rollback()
        raise

@router.patch("/{lkh_id}", response_model=LKHResponse)
def update_lkh(
    lkh_id: int,
    data: LKHUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(
        require_roles("SUPERADMIN", "SPV", "SALES", "PPIC")
    ),
):
    laporan = get_lkh_query(db, current_user).filter(
        LKH.id == lkh_id
    ).first()

    if not laporan:
        raise HTTPException(
            404,
            "Laporan tidak ditemukan atau kamu tidak memiliki akses",
        )

    updates = data.model_dump(exclude_unset=True)

    for field, value in updates.items():
        setattr(laporan, field, value)

    if "status_lkh" in updates:
        rkh = db.query(RKH).filter(RKH.id == laporan.rkh_id).first()
        if rkh:
            rkh.status = "CLOSED" if laporan.status_lkh == "DONE" else "OPEN"

    try:
        db.commit()
        db.refresh(laporan)
        return laporan
    except Exception:
        db.rollback()
        raise

@router.delete("/{lkh_id}", status_code=204)
def delete_lkh(
    lkh_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(
        require_roles("SUPERADMIN", "SPV", "SALES")
    ),
):
    laporan = get_lkh_query(db, current_user).filter(
        LKH.id == lkh_id
    ).first()

    if not laporan:
        raise HTTPException(
            404,
            "Laporan tidak ditemukan atau kamu tidak memiliki akses",
        )

    db.delete(laporan)
    db.commit()

    return Response(status_code=204)