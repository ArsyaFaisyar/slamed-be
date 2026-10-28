from fastapi import APIRouter, Depends, HTTPException, Query, Response
from sqlalchemy import func
from sqlalchemy.orm import Session

from pydantic import BaseModel, Field
from io import BytesIO
from datetime import date, time, timedelta

from fastapi.responses import StreamingResponse
from sqlalchemy.orm import joinedload

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.db.database import get_db
from app.models.customer import Customer
from app.models.user import User
from app.models.area import Area
from app.schemas.customer import (
    CustomerCreate,
    CustomerUpdate,
    CustomerResponse,
)
from app.core.dependencies import get_current_user, require_roles
from app.core.roles import CUSTOMER_VIEW_ROLES
from datetime import datetime, timezone
from zoneinfo import ZoneInfo

router = APIRouter(
    prefix="/customers",
    tags=["Customers"],
)


def validate_area(db: Session, area_id: int | None):
    if area_id is None:
        return

    area = db.query(Area).filter(
        Area.id == area_id
    ).first()

    if not area:
        raise HTTPException(
            status_code=400,
            detail="Area tidak ditemukan",
        )


def validate_sales(db: Session, sales_id: int | None):
    if sales_id is None:
        return

    sales = db.query(User).filter(
        User.id == sales_id,
        User.role == "SALES",
        User.is_active == 1,
    ).first()

    if not sales:
        raise HTTPException(
            status_code=400,
            detail="Sales tidak ditemukan, tidak aktif, atau role bukan SALES",
        )


WIB = ZoneInfo("Asia/Jakarta")

STATUS_APPROVAL_LABELS = {
    "PENDING": "Menunggu Persetujuan",
    "APPROVED": "Disetujui",
    "REJECTED": "Ditolak",
}

STATUS_APPROVAL_STYLES = {
    "PENDING": ("FFF2CC", "7F6000"),
    "APPROVED": ("E2EFDA", "375623"),
    "REJECTED": ("F8CBAD", "9C0006"),
}


def to_wib(value):
    if value is None:
        return None
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.astimezone(WIB).replace(tzinfo=None)


def clean_values(values):
    return [item.strip() for item in (values or []) if item and item.strip()]


def build_customer_workbook(customers, filter_text, exported_by, exported_at):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Data Customer"

    headers = [
        "Nama Instansi",
        "Nama PIC",
        "Kontak PIC",
        "Kota",
        "Kategori",
        "Area",
        "Sales",
        "Status Approval",
        "Alasan Penolakan",
        "Tanggal Dibuat",
        "Terakhir Diperbarui",
    ]

    worksheet.append(headers)

    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78",
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    for customer in customers:
        status = customer.approval_status

        worksheet.append([
            customer.nama_instansi,
            customer.nama_pic,
            str(customer.kontak_pic) if customer.kontak_pic is not None else None,
            customer.kota,
            customer.kategori,
            customer.area_nama,
            customer.sales_nama,
            STATUS_APPROVAL_LABELS.get(status, status),
            customer.rejection_reason,
            to_wib(customer.created_at),
            to_wib(customer.updated_at),
        ])

    for row in worksheet.iter_rows(min_row=2):
        row[2].number_format = "@"
        row[9].number_format = "dd-mm-yyyy hh:mm:ss"
        row[10].number_format = "dd-mm-yyyy hh:mm:ss"

    for column_cells in worksheet.columns:
        max_length = max(
            len(str(cell.value or ""))
            for cell in column_cells
        )
        column_letter = get_column_letter(
            column_cells[0].column
        )
        worksheet.column_dimensions[
            column_letter
        ].width = min(max_length + 3, 40)

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    return workbook


@router.get("/", response_model=list[CustomerResponse])
def get_customers(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    query = db.query(Customer)

    if current_user.role == "SALES":
        query = query.filter(Customer.sales_id == current_user.id)
    elif current_user.role not in CUSTOMER_VIEW_ROLES:
        raise HTTPException(
            status_code=403,
            detail="Kamu tidak memiliki akses ke data customer",
        )

    return query.order_by(Customer.id.desc()).all()

@router.get("/export/excel")
def export_customer_excel(
    area: list[str] | None = Query(None),
    sales: list[str] | None = Query(None),
    status_approval: list[str] | None = Query(None),
    tanggal_mulai: date | None = None,
    tanggal_selesai: date | None = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    if tanggal_mulai and tanggal_selesai:
        if tanggal_mulai > tanggal_selesai:
            raise HTTPException(
                status_code=400,
                detail="Tanggal mulai tidak boleh melebihi tanggal selesai",
            )

    query = db.query(Customer).options(
        joinedload(Customer.area_rel),
        joinedload(Customer.sales_rel),
    )

    if current_user.role == "SALES":
        query = query.filter(Customer.sales_id == current_user.id)
    elif current_user.role not in CUSTOMER_VIEW_ROLES:
        raise HTTPException(
            status_code=403,
            detail="Kamu tidak memiliki akses ke data customer",
        )

    areas = clean_values(area)
    sales_names = clean_values(sales)
    statuses = [item.upper() for item in clean_values(status_approval)]

    if areas:
        query = query.filter(
            Customer.area_rel.has(
                func.lower(Area.nama).in_([item.lower() for item in areas])
            )
        )

    if sales_names:
        query = query.filter(
            Customer.sales_rel.has(
                func.lower(User.nama).in_([item.lower() for item in sales_names])
            )
        )

    if statuses:
        query = query.filter(Customer.approval_status.in_(statuses))

    if tanggal_mulai:
        query = query.filter(
            Customer.created_at
            >= datetime.combine(tanggal_mulai, time.min, tzinfo=WIB)
        )

    if tanggal_selesai:
        query = query.filter(
            Customer.created_at
            < datetime.combine(
                tanggal_selesai + timedelta(days=1), time.min, tzinfo=WIB
            )
        )

    data_customer = query.order_by(Customer.id.desc()).all()

    def join_or_all(values, empty="Semua"):
        return ", ".join(values) if values else empty

    if tanggal_mulai or tanggal_selesai:
        mulai_text = tanggal_mulai.strftime("%d-%m-%Y") if tanggal_mulai else "..."
        selesai_text = tanggal_selesai.strftime("%d-%m-%Y") if tanggal_selesai else "..."
        periode_text = f"{mulai_text} s/d {selesai_text}"
    else:
        periode_text = "Semua"

    filter_text = (
        f"Area: {join_or_all(areas)} | "
        f"Sales: {join_or_all(sales_names)} | "
        f"Status: {join_or_all([STATUS_APPROVAL_LABELS.get(item, item) for item in statuses])} | "
        f"Periode dibuat: {periode_text}"
    )

    now = datetime.now(WIB)
    workbook = build_customer_workbook(
        data_customer,
        filter_text,
        current_user.nama,
        now,
    )

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    filename = f"Data_Customer_{now.strftime('%d-%m-%Y_%H%M')}.xlsx"

    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"'
        },
    )


@router.get("/{customer_id}", response_model=CustomerResponse)
def get_customer(
    customer_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    customer = db.query(Customer).filter(
        Customer.id == customer_id
    ).first()

    if not customer:
        raise HTTPException(
            status_code=404,
            detail="Customer tidak ditemukan",
        )

    if current_user.role == "SALES":
        if customer.sales_id != current_user.id:
            raise HTTPException(
                status_code=403,
                detail="Customer ini bukan tanggung jawabmu",
            )
    elif current_user.role not in CUSTOMER_VIEW_ROLES:
        raise HTTPException(
            status_code=403,
            detail="Kamu tidak memiliki akses ke data customer",
        )

    return customer


@router.post("/", response_model=CustomerResponse, status_code=201)
def create_customer(
    data: CustomerCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(
        require_roles("SALES", "SUPERADMIN", "SPV")
    ),
):
    if current_user.role == "SALES":
        if current_user.area_id is None:
            raise HTTPException(
                status_code=400,
                detail="Akun Anda belum memiliki area",
            )

        customer = Customer(
            **data.model_dump(exclude={"sales_id", "area_id"}),
            sales_id=current_user.id,
            area_id=current_user.area_id,
            approval_status="PENDING",
        )
    else:
        validate_sales(db, data.sales_id)
        validate_area(db, data.area_id)

        customer = Customer(
            **data.model_dump(),
            approval_status="APPROVED",
            verified_by=current_user.id,
            verified_at=datetime.now(timezone.utc),
        )

    db.add(customer)
    db.commit()
    db.refresh(customer)

    return customer


@router.patch("/{customer_id}", response_model=CustomerResponse)
def update_customer(
    customer_id: int,
    data: CustomerUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(
        require_roles("SALES", "SPV", "SUPERADMIN")
    ),
):
    customer = db.query(Customer).filter(
        Customer.id == customer_id
    ).first()

    if not customer:
        raise HTTPException(
            status_code=404,
            detail="Customer tidak ditemukan",
        )

    updates = data.model_dump(exclude_unset=True)

    if current_user.role == "SALES":
        if customer.sales_id != current_user.id:
            raise HTTPException(
                status_code=403,
                detail="Customer ini bukan tanggung jawabmu",
            )

        if customer.approval_status == "APPROVED":
            raise HTTPException(
                status_code=403,
                detail="Customer yang sudah disetujui tidak bisa diedit oleh Sales",
            )

        updates.pop("sales_id", None)
        updates.pop("area_id", None)

    else:
        if "area_id" in updates:
            validate_area(db, updates["area_id"])

        if "sales_id" in updates:
            validate_sales(db, updates["sales_id"])

    for field, value in updates.items():
        setattr(customer, field, value)

    if current_user.role == "SALES":
        customer.approval_status = "PENDING"
        customer.verified_by = None
        customer.verified_at = None
        customer.rejection_reason = None

    db.commit()
    db.refresh(customer)

    return customer


@router.patch("/{customer_id}/approve", response_model=CustomerResponse)
def approve_customer(
    customer_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(
        require_roles("SPV", "SUPERADMIN")
    ),
):
    customer = db.query(Customer).filter(
        Customer.id == customer_id
    ).first()

    if not customer:
        raise HTTPException(
            status_code=404,
            detail="Customer tidak ditemukan",
        )

    if customer.approval_status == "APPROVED":
        raise HTTPException(
            status_code=400,
            detail="Customer sudah disetujui",
        )

    customer.approval_status = "APPROVED"
    customer.verified_by = current_user.id
    customer.verified_at = datetime.now(timezone.utc)
    customer.rejection_reason = None

    db.commit()
    db.refresh(customer)

    return customer


@router.delete("/{customer_id}", status_code=204)
def delete_customer(
    customer_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(
        require_roles("SUPERADMIN", "SPV")
    ),
):
    customer = db.query(Customer).filter(
        Customer.id == customer_id
    ).first()

    if not customer:
        raise HTTPException(
            status_code=404,
            detail="Customer tidak ditemukan",
        )

    db.delete(customer)
    db.commit()

    return Response(status_code=204)


class CustomerRejectRequest(BaseModel):
    alasan: str = Field(min_length=1, max_length=1000)


@router.patch("/{customer_id}/reject", response_model=CustomerResponse)
def reject_customer(
    customer_id: int,
    data: CustomerRejectRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(
        require_roles("SPV", "SUPERADMIN")
    ),
):
    alasan = data.alasan.strip()

    if not alasan:
        raise HTTPException(
            status_code=422,
            detail="Alasan penolakan wajib diisi",
        )

    customer = db.query(Customer).filter(
        Customer.id == customer_id
    ).first()

    if not customer:
        raise HTTPException(
            status_code=404,
            detail="Customer tidak ditemukan",
        )

    if customer.approval_status == "APPROVED":
        raise HTTPException(
            status_code=400,
            detail="Customer yang sudah disetujui tidak bisa ditolak",
        )

    customer.approval_status = "REJECTED"
    customer.verified_by = current_user.id
    customer.verified_at = datetime.now(timezone.utc)
    customer.rejection_reason = alasan

    db.commit()
    db.refresh(customer)

    return customer